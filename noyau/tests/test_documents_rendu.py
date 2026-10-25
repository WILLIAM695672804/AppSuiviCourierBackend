"""
Le rendu Word — la seule couche qui connait python-docx.

Ce que ces tests protegent :

 1. **L'en-tete se repete.** C'est le defaut qu'on ne voit qu'a la troisieme
    page, quand la note est deja diffusee.
 2. **Aucune mise en forme directe.** Le rendu passe par les styles nommes du
    gabarit ; une taille de police posee dans le code diverge du modele des
    que la DSR remet un nouveau fichier.
 3. **La numerotation rendue est celle que la structure a calculee.**
"""

import docx
import pytest

from noyau.documents import rendu_docx
from noyau.documents.structure import (
    Cellule,
    Chiffre,
    Chiffres,
    Colonne,
    Document,
    Encadre,
    Paragraphe,
    Puces,
    Section,
    Tableau,
)


@pytest.fixture
def document() -> Document:
    return Document(
        titre="Note de synthèse — Comité de direction",
        surtitre="Division du suivi et de la relance",
        sous_titre="Séance du 3 septembre 2026 · période observée : août 2026",
        reference="NOTE-CODIR-2026-09",
        sections=(
            Section(
                titre="Faits marquants de la période",
                blocs=(Puces(items=("Le taux d'exécution s'établit à 62 %.",)),),
            ),
            Section(
                titre="Chiffres d'ensemble",
                blocs=(Chiffres(items=(Chiffre(libelle="Exécution", valeur="62 %"),)),),
            ),
            Section(
                titre="Classement des structures",
                blocs=(
                    Tableau(
                        colonnes=(
                            Colonne(entete="Rang"),
                            Colonne(entete="Structure", poids=4),
                            Colonne(entete="Score", a_droite=True),
                            Colonne(entete="Cotation", poids=2),
                        ),
                        lignes=(
                            (
                                Cellule(texte="1"),
                                Cellule(texte="DSR — Division du suivi"),
                                Cellule(texte="88"),
                                Cellule(texte="Vert", barre=88, feu="vert"),
                            ),
                        ),
                        legende="Score sur 100, au dernier mois clos.",
                    ),
                ),
            ),
            Section(
                titre="Points appelant un arbitrage",
                blocs=(
                    Encadre(
                        titre="Décisions échues",
                        texte="Trois restent à exécuter.",
                        feu="rouge",
                    ),
                ),
            ),
        ),
        annexes=(
            Section(
                titre="Classement complet",
                blocs=(Paragraphe(texte="Le détail des 24 structures."),),
            ),
        ),
    )


def test_le_fichier_est_ecrit(tmp_path, document):
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")

    assert chemin.exists()
    assert chemin.stat().st_size > 0


def test_l_entete_institutionnel_se_repete_sur_chaque_page(tmp_path, document):
    """
    Le gabarit porte l'en-tete dans `header1.xml`, pas dans le corps.

    C'est le piege du modele remis par la DSR : recopie tel quel, il ne sort
    qu'en page 1.
    """
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    entete = rendu.sections[0].header
    assert entete.is_linked_to_previous is False
    assert len(entete.tables) == 1
    # Le modele de la DSR ecrit l'en-tete en capitales : « RÉPUBLIQUE DU
    # CAMEROUN ». On compare sans la casse plutot que de la corriger — c'est
    # la piece qui fait foi.
    assert "CAMEROUN" in entete.tables[0].cell(0, 0).text.upper()


def test_la_colonne_anglaise_de_l_entete_est_la(tmp_path, document):
    """Le bilinguisme est la forme officielle des actes : il ne se perd pas."""
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    texte = "\n".join(cellule.text for cellule in rendu.sections[0].header.tables[0]._cells)
    assert "REPUBLIC OF CAMEROON" in texte.upper()
    assert "MONITORING AND RECOVERY DIVISION" in texte.upper()


def test_les_sections_sortent_numerotees(tmp_path, document):
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    titres = [
        paragraphe.text
        for paragraphe in rendu.paragraphs
        if paragraphe.style.name == "Titre section"
    ]
    assert titres == [
        "1. Faits marquants de la période",
        "2. Chiffres d'ensemble",
        "3. Classement des structures",
        "4. Points appelant un arbitrage",
    ]


def test_l_annexe_sort_lettree(tmp_path, document):
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    titres = [
        paragraphe.text
        for paragraphe in rendu.paragraphs
        if paragraphe.style.name == "Titre annexe"
    ]
    assert titres == ["Annexe A — Classement complet"]


def test_le_rendu_n_ecrit_aucune_mise_en_forme_directe(tmp_path, document):
    """
    Tout passe par un style nomme du gabarit.

    Une taille posee a la main survit au premier rendu et diverge du modele au
    suivant. Ce test est le seul garde-fou : rien d'autre ne le signalerait.
    """
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    for paragraphe in rendu.paragraphs:
        for passage in paragraphe.runs:
            assert passage.font.size is None, f"Taille en dur : « {paragraphe.text} »"
            assert passage.font.bold is None, f"Gras en dur : « {paragraphe.text} »"


def test_aucune_mise_en_forme_directe_dans_les_cellules(tmp_path, document):
    """
    Le garde-fou vaut aussi sous les tableaux.

    `rendu.paragraphs` ne descend pas dans les cellules : la moitie du document
    — chiffres, classements, encadres — echapperait au test precedent.
    """
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    for tableau in rendu.tables:
        for cellule in tableau._cells:
            for paragraphe in cellule.paragraphs:
                for passage in paragraphe.runs:
                    assert passage.font.size is None, f"Taille en dur : « {paragraphe.text} »"
                    assert passage.font.bold is None, f"Gras en dur : « {paragraphe.text} »"


def test_le_tableau_porte_ses_en_tetes_et_ses_lignes(tmp_path, document):
    # `rendu.tables[0]` est la rangee de chiffres de la section 2 : elle est
    # rendue en tableau sans filet, et elle vient avant le classement.
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    tableau = rendu.tables[1]
    assert [cellule.text for cellule in tableau.rows[0].cells] == [
        "Rang",
        "Structure",
        "Score",
        "Cotation",
    ]
    assert tableau.rows[1].cells[1].text == "DSR — Division du suivi"


def test_la_reference_est_posee_dans_le_pied(tmp_path, document):
    """C'est ce qu'on cite quand on rappelle la note en seance."""
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    pied = "\n".join(paragraphe.text for paragraphe in rendu.sections[0].footer.paragraphs)
    assert "NOTE-CODIR-2026-09" in pied


def test_la_reference_n_efface_ni_le_bandeau_ni_la_pagination(tmp_path, document):
    """
    Le pied du gabarit porte deja deux choses, et elles restent.

    Le bandeau legal est une **image ancree** dans le premier paragraphe du
    pied : ecrire la reference par-dessus ce paragraphe la ferait disparaitre
    sans qu'aucun test ne bronche. La « Page X sur Y » est dans le second.
    """
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    pied = rendu.sections[0].footer
    xml = pied._element.xml
    assert "Pied_PAD" in xml, "Le bandeau legal a disparu du pied."
    assert 'w:instr=" PAGE "' in xml
    assert 'w:instr=" NUMPAGES "' in xml
    assert "NOTE-CODIR-2026-09" in "\n".join(p.text for p in pied.paragraphs)


def test_la_barre_de_score_est_faite_de_cellules_ombrees(tmp_path, document):
    """
    Pas une image : le secretariat corrige la note avant la seance (ADR 023).

    La barre est un tableau imbrique de deux cellules teintees, dans la cellule
    de cotation du classement.
    """
    chemin = rendu_docx.ecrire(document, tmp_path / "note.docx")
    rendu = docx.Document(str(chemin))

    cotation = rendu.tables[1].rows[1].cells[3]
    assert len(cotation.tables) == 1
    assert len(cotation.tables[0].rows[0].cells) == 2
    assert rendu_docx.BARRE_PLEINE in cotation.tables[0].rows[0].cells[0]._tc.xml
    assert "graphicData" not in cotation._tc.xml, "Une image s'est glissée dans la barre."
