"""
Fabrique le gabarit des documents depuis le modele remis par la DSR.

    ./.venv/bin/python scripts/preparer_gabarit.py

**A lancer a la main, jamais au demarrage.** Le resultat est versionne : un
gabarit refabrique a chaque deploiement serait un gabarit dont personne ne peut
dire ce qu'il contient.

Ce que le script corrige, et pourquoi il existe :

 1. **L'en-tete du modele est un tableau pose dans le corps**, pas un en-tete de
    page. Recopie tel quel, il n'apparaitrait qu'en page 1 — et une note de
    synthese de trois pages sortirait avec deux pages anonymes sous la
    signature du Directeur general. Le script le deplace dans `header1.xml`.

 2. **Deplacer un tableau qui porte des images casse leurs liens.** Les
    `r:embed` d'un `w:tbl` pointent vers les relations de `document.xml` ;
    l'en-tete de page a les siennes. Sans re-liaison, Word et LibreOffice
    affichent un cadre vide a la place des armoiries. Le script relie chaque
    image a la partie d'en-tete et reecrit les identifiants.

 3. Le pied de page du modele — le bandeau legal — est deja un vrai pied et se
    repete. Le script y ajoute « Page X sur Y ».

 4. Les styles nommes que `rendu_docx` emploie sont crees ici. **Le rendu
    n'ecrit jamais de mise en forme directe** : une taille de police posee a la
    main dans le code diverge du modele au premier changement.

 5. **La marge haute est recalculee** depuis la hauteur mesuree de l'en-tete.
    Le modele portait l'en-tete dans le corps : il n'avait rien a reserver en
    haut de page. Une fois l'en-tete devenu un vrai en-tete, la marge doit lui
    faire place. Voir `_poser_la_marge_haute`.
"""

from copy import deepcopy
from pathlib import Path
from typing import Any

import docx
from docx.document import Document as DocumentWord
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as TYPE_RELATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Length, Pt, RGBColor, Twips

RACINE = Path(__file__).resolve().parent.parent
MODELE = RACINE.parent / "frontend" / "public" / "Entête DSR à utiliser.docx"
GABARIT = RACINE / "noyau" / "documents" / "gabarits" / "pad-officiel.docx"

#: Le gris de l'institution, repris de la charte du front (`--color-encre`).
ENCRE = RGBColor(0x1A, 0x1F, 0x2B)
ENCRE_TENUE = RGBColor(0x66, 0x70, 0x85)


def preparer() -> Path:
    if not MODELE.exists():
        raise SystemExit(f"Modele introuvable : {MODELE}")

    document = docx.Document(str(MODELE))
    _deplacer_entete_dans_la_page(document)
    _poser_la_marge_haute(document)
    _vider_le_corps(document)
    _completer_le_pied(document)
    _poser_les_styles(document)

    GABARIT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(GABARIT))
    return GABARIT


def _deplacer_entete_dans_la_page(document: DocumentWord) -> None:
    """Le tableau d'en-tete passe du corps a `header1.xml`, images comprises."""
    if len(document.tables) != 1:
        raise SystemExit(
            f"Le modele porte {len(document.tables)} tableaux, un seul etait attendu : "
            "l'en-tete. Verifier que le fichier remis par la DSR n'a pas change."
        )

    section = document.sections[0]
    # Sans cela, python-docx ne cree pas `header1.xml` : l'en-tete resterait
    # herite de la section precedente, c'est-a-dire vide.
    section.header.is_linked_to_previous = False

    tableau = deepcopy(document.tables[0]._element)
    _relier_les_images(tableau, document.part, section.header.part)

    # L'en-tete d'une section neuve porte un paragraphe vide : place avant le
    # tableau, il le pousserait d'une ligne sur chaque page. On le vide de sa
    # mise en forme et on le rejette apres le tableau, ou il sert de
    # paragraphe de cloture — un `w:tbl` ne doit pas terminer une partie
    # d'en-tete, Word reecrit le fichier sinon.
    for paragraphe in list(section.header.paragraphs):
        paragraphe._element.getparent().remove(paragraphe._element)

    section.header._element.append(tableau)
    cloture = section.header.add_paragraph()
    cloture.paragraph_format.space_before = Pt(0)
    cloture.paragraph_format.space_after = Pt(0)
    cloture.add_run().font.size = Pt(1)


def _relier_les_images(element: Any, partie_source: Any, partie_cible: Any) -> None:
    """
    Reecrit les `r:embed` du fragment copie vers les relations de la cible.

    C'est le point 2 du docstring du module. Sans lui, le gabarit s'ouvre sans
    erreur et affiche un cadre vide — un defaut qu'on ne voit qu'a
    l'impression, quand la note est deja partie.
    """
    for blip in element.iter(qn("a:blip")):
        ancien = blip.get(qn("r:embed"))
        if ancien is None:
            continue
        image = partie_source.related_parts[ancien]
        blip.set(qn("r:embed"), partie_cible.relate_to(image, TYPE_RELATION.IMAGE))


#: Hauteur d'une ligne a interligne simple, en multiple de la taille de police.
#: Word et LibreOffice la tirent des metriques de la fonte ; 1,2 est la valeur
#: courante et la plus haute des deux moteurs sur les fontes usuelles.
FACTEUR_INTERLIGNE = 1.2

#: Taille de police retenue quand un passage n'en declare aucune : 11 pt, le
#: defaut de Word, exprime en demi-points comme dans `w:sz`.
TAILLE_PAR_DEFAUT = 22

#: Respiration entre le bas de l'en-tete et la premiere ligne du corps. Elle
#: absorbe aussi ce que la mesure ne sait pas voir : un libelle qui passerait
#: sur deux lignes selon la fonte reellement installee.
RESPIRATION_SOUS_ENTETE = Cm(0.5)


def _poser_la_marge_haute(document: DocumentWord) -> None:
    """
    Reserve en haut de page la place que l'en-tete occupe desormais.

    Le modele de la DSR portait l'en-tete **dans le corps** : il n'avait rien a
    reserver, et sa marge haute de 426 twips (0,75 cm) etait coherente. En
    deplacant l'en-tete dans un vrai en-tete de page (point 1 du docstring du
    module), la marge haute doit maintenant contenir la distance d'en-tete
    **plus** la hauteur de l'en-tete lui-meme. Faute de quoi la zone d'en-tete
    deborde la marge, et Word et LibreOffice repoussent le corps chacun a sa
    facon : le meme gabarit ne rend pas la meme page selon le logiciel.

    La valeur n'est pas choisie, elle est **calculee** depuis la hauteur mesuree
    de l'en-tete — cf. `_hauteur_de_l_entete`. Les autres marges restent celles
    du modele : on ne touche qu'a ce que le deplacement a rendu faux.
    """
    section = document.sections[0]
    distance = section.header_distance
    if distance is None:
        raise SystemExit(
            "Le modele ne declare pas de distance d'en-tete (`w:header`) : impossible de "
            "calculer la marge haute. Verifier le fichier remis par la DSR."
        )
    section.top_margin = Emu(
        distance + _hauteur_de_l_entete(section.header) + RESPIRATION_SOUS_ENTETE
    )


def _hauteur_de_l_entete(entete: Any) -> Length:
    """
    Hauteur de l'en-tete, mesuree depuis son contenu.

    Le tableau d'en-tete a trois lignes et trois colonnes : deux colonnes de
    texte bilingue, et au centre une colonne fusionnee (`w:vMerge`) qui ne
    porte que les armoiries. Les deux ne s'additionnent pas — c'est la plus
    haute des deux qui commande. On y ajoute le paragraphe de cloture pose par
    `_deplacer_entete_dans_la_page`.
    """
    tableau = entete.tables[0]._tbl

    hauteur_texte = 0
    for ligne in tableau.findall(qn("w:tr")):
        cellules_de_texte = [
            cellule
            for cellule in ligne.findall(qn("w:tc"))
            if cellule.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}") is None
        ]
        hauteur_texte += max(
            (_hauteur_du_texte(cellule) for cellule in cellules_de_texte),
            default=0,
        )

    # `wp:extent` donne la hauteur d'**affichage** de l'image, en EMU : c'est
    # celle qui occupe la page, pas celle du fichier PNG.
    hauteur_images = max(
        (int(extent.get("cy")) for extent in tableau.iter(qn("wp:extent"))),
        default=0,
    )

    cloture = sum(_hauteur_du_paragraphe(p._p) for p in entete.paragraphs)
    return Emu(max(hauteur_texte, hauteur_images) + cloture)


def _hauteur_du_texte(cellule: Any) -> int:
    """Hauteur en EMU du texte d'une cellule, une ligne par paragraphe."""
    return sum(_hauteur_du_paragraphe(p) for p in cellule.findall(qn("w:p")))


def _hauteur_du_paragraphe(paragraphe: Any) -> int:
    """
    Hauteur en EMU d'un paragraphe d'une seule ligne.

    Tous les paragraphes de l'en-tete sont a interligne simple (`w:line="240"`,
    `w:lineRule="auto"`) : la hauteur de ligne vaut la plus grande taille de
    police du paragraphe, majoree du facteur d'interligne, plus les espacements
    avant et apres.
    """
    demi_points = [
        int(taille.get(qn("w:val")))
        for taille in paragraphe.iter(qn("w:sz"))
        if taille.get(qn("w:val")) is not None
    ]
    ligne = Pt(max(demi_points, default=TAILLE_PAR_DEFAUT) / 2) * FACTEUR_INTERLIGNE

    espacement = paragraphe.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
    avant = apres = 0
    if espacement is not None:
        avant = Twips(int(espacement.get(qn("w:before")) or 0))
        apres = Twips(int(espacement.get(qn("w:after")) or 0))
    return int(ligne) + avant + apres


def _vider_le_corps(document: DocumentWord) -> None:
    """
    Le corps est vide : c'est `rendu_docx` qui l'ecrit.

    `w:sectPr` est conserve — il porte le format de page, les marges et les
    references d'en-tete et de pied. Le supprimer rendrait un document sans
    en-tete du tout.
    """
    corps = document.element.body
    for enfant in list(corps):
        if enfant.tag != qn("w:sectPr"):
            corps.remove(enfant)


def _completer_le_pied(document: DocumentWord) -> None:
    """
    « Page X sur Y » sous le bandeau legal.

    Les champs `PAGE` et `NUMPAGES` sont ecrits en `w:fldSimple` : c'est la
    forme que Word et LibreOffice lisent tous les deux. Un numero calcule au
    rendu aurait ete faux des que la pagination bouge d'une ligne.
    """
    pied = document.sections[0].footer
    pied.is_linked_to_previous = False

    ligne = pied.add_paragraph()
    ligne.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def _texte(contenu: str) -> None:
        passage = ligne.add_run(contenu)
        passage.font.size = Pt(8)
        passage.font.color.rgb = ENCRE_TENUE

    _texte("Page ")
    ligne._p.append(_champ("PAGE"))
    _texte(" sur ")
    ligne._p.append(_champ("NUMPAGES"))


def _champ(instruction: str) -> Any:
    """Un champ Word simple — `PAGE`, `NUMPAGES`."""
    champ = OxmlElement("w:fldSimple")
    champ.set(qn("w:instr"), f" {instruction} ")
    passage = OxmlElement("w:r")
    texte = OxmlElement("w:t")
    texte.text = "1"
    passage.append(texte)
    champ.append(passage)
    return champ


#: Les styles que `rendu_docx` emploie, et rien de plus. Chaque entree :
#: (nom, taille en points, gras, couleur, espace avant, espace apres).
STYLES = (
    ("Titre note", 17, True, ENCRE, 0, 4),
    ("Sous-titre note", 9.5, False, ENCRE_TENUE, 0, 14),
    ("Surtitre note", 8.5, True, ENCRE_TENUE, 0, 3),
    ("Titre section", 11, True, ENCRE, 14, 6),
    ("Corps document", 10, False, ENCRE, 0, 5),
    ("Puce document", 10, False, ENCRE, 0, 3),
    ("Libelle chiffre", 8, True, ENCRE_TENUE, 0, 1),
    ("Valeur chiffre", 16, True, ENCRE, 0, 6),
    ("Entete tableau", 8, True, ENCRE_TENUE, 2, 2),
    ("Cellule tableau", 9, False, ENCRE, 2, 2),
    ("Legende tableau", 8, False, ENCRE_TENUE, 4, 0),
    ("Titre encadre", 10, True, ENCRE, 8, 2),
    ("Titre annexe", 12, True, ENCRE, 16, 6),
)


#: Le style de tableau qu'emploie `rendu_docx` pour les classements et les
#: encadres. Il n'existe pas dans le modele de la DSR, et « Table Grid » n'est
#: pas davantage dans un document francais : Word ne cree ses styles latents
#: qu'a l'usage. Le poser ici plutot que d'ecrire les bordures au rendu tient
#: la regle du chantier — le trait est une affaire de gabarit, pas de code.
STYLE_TABLEAU = "Tableau document"

#: Le gris des filets. Assez tenu pour ne pas manger la lecture, assez present
#: pour survivre a une photocopie.
FILET = "C9CFDA"

#: Epaisseur des filets en huitiemes de point (`w:sz`) : 4 vaut un demi-point.
FILET_EPAISSEUR = "4"


def _poser_les_styles(document: DocumentWord) -> None:
    for nom, taille, gras, couleur, avant, apres in STYLES:
        if nom in [style.name for style in document.styles]:
            continue
        style = document.styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.size = Pt(taille)
        style.font.bold = gras
        style.font.color.rgb = couleur
        style.paragraph_format.space_before = Pt(avant)
        style.paragraph_format.space_after = Pt(apres)

    _poser_le_style_de_tableau(document)


def _poser_le_style_de_tableau(document: DocumentWord) -> None:
    """
    Un style de tableau a filets, ecrit en XML.

    python-docx n'expose pas les bordures d'un style de tableau : il faut poser
    `w:tblPr/w:tblBorders` a la main. C'est le prix a payer pour que le rendu
    reste sans mise en forme directe — le trait vit dans le gabarit, et une DSR
    qui remet un modele le remet avec ses filets.

    `w:tblPr` se place apres `w:rPr` dans un `w:style` : c'est l'ordre du
    schema, et Word refuse d'ouvrir un fichier qui s'en ecarte.
    """
    if STYLE_TABLEAU in [style.name for style in document.styles]:
        return

    style = document.styles.add_style(STYLE_TABLEAU, WD_STYLE_TYPE.TABLE)
    style.base_style = document.styles["Normal Table"]

    proprietes = OxmlElement("w:tblPr")
    bordures = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        filet = OxmlElement(f"w:{cote}")
        filet.set(qn("w:val"), "single")
        filet.set(qn("w:sz"), FILET_EPAISSEUR)
        filet.set(qn("w:space"), "0")
        filet.set(qn("w:color"), FILET)
        bordures.append(filet)
    proprietes.append(bordures)
    style.element.append(proprietes)


if __name__ == "__main__":
    print(f"Gabarit ecrit : {preparer()}")
